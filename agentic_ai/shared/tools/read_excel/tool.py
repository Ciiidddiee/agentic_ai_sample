"""Shared Excel extractor (format plumbing for any agent)."""

from __future__ import annotations

from pathlib import Path
from typing import Any


def read_excel(document_path: Path) -> dict[str, Any]:
    """Extract sheet contents as searchable text.

    Returns: text, confidence, kind, sheets, sheet_count
    """
    from openpyxl import load_workbook

    path = Path(document_path)
    if not path.is_file():
        raise FileNotFoundError(f"Document not found: {path}")

    wb = load_workbook(path, data_only=True, read_only=True)
    sheets_out: list[dict[str, Any]] = []
    parts: list[str] = []

    for sheet in wb.worksheets:
        rows: list[str] = []
        for row in sheet.iter_rows(values_only=True):
            cells = ["" if c is None else str(c) for c in row]
            if any(cell.strip() for cell in cells):
                rows.append("\t".join(cells))
        sheet_text = "\n".join(rows)
        sheets_out.append(
            {"name": sheet.title, "text": sheet_text, "row_count": len(rows)}
        )
        if sheet_text.strip():
            parts.append(f"## Sheet: {sheet.title}\n{sheet_text}")

    wb.close()
    full_text = "\n\n".join(parts).strip()
    confidence = 1.0 if full_text else 0.0

    return {
        "text": full_text,
        "confidence": confidence,
        "kind": "excel",
        "sheets": sheets_out,
        "sheet_count": len(sheets_out),
    }
